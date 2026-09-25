"""Проверка ворот: чистая колода проходит, испорченная — ловится по каждой поломке.

Смысл не в том, что «скрипт запустился», а в том, что он РАЗЛИЧАЕТ. Колода
портится нарочно четырьмя способами, и ворота обязаны назвать каждый.

    python test-deck-gates.py
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.util import Cm, Pt

HERE = Path(__file__).resolve().parent
SCRIPT = HERE / "deck_gates.py"


def blank_deck(path: Path, *, broken: bool) -> None:
    deck = Presentation()
    deck.slide_width, deck.slide_height = Cm(33.87), Cm(19.05)

    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Заголовок"
    for paragraph in slide.shapes.title.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(30)
    body = slide.shapes.add_textbox(Cm(2), Cm(5), Cm(20), Cm(4))
    body.text_frame.text = "Обычный текст на слайде, кегль нормальный."
    for paragraph in body.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(14)

    if broken:
        # 1. слишком мелкий шрифт
        tiny = slide.shapes.add_textbox(Cm(2), Cm(10), Cm(10), Cm(2))
        tiny.text_frame.text = "Мелкий текст"
        for paragraph in tiny.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(6)

        # 2. фигура за краем слайда
        outside = slide.shapes.add_textbox(Cm(-4), Cm(12), Cm(8), Cm(2))
        outside.text_frame.text = "Уехало влево"
        for paragraph in outside.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(14)

        # 3. остаток-заглушка
        todo = slide.shapes.add_textbox(Cm(20), Cm(14), Cm(10), Cm(2))
        todo.text_frame.text = "TODO: дописать"
        for paragraph in todo.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(14)

        # 4. таблица с неверным «Итого»
        table = slide.shapes.add_table(3, 2, Cm(2), Cm(15), Cm(12), Cm(3)).table
        table.cell(0, 0).text = "статья"
        table.cell(0, 1).text = "сумма"
        table.cell(1, 0).text = "первое"
        table.cell(1, 1).text = "100"
        table.cell(2, 0).text = "Итого"
        table.cell(2, 1).text = "250"

    deck.save(str(path))


def run(deck: Path) -> tuple[int, dict]:
    result = subprocess.run([sys.executable, str(SCRIPT), str(deck), "--json"],
                            capture_output=True, text=True, encoding="utf-8")
    return result.returncode, json.loads(result.stdout)


def main() -> int:
    work = Path(tempfile.mkdtemp(prefix="deck-gates-test-"))
    try:
        clean = work / "clean.pptx"
        blank_deck(clean, broken=False)
        code, report = run(clean)
        assert code == 0, f"чистая колода не прошла: {report['findings']}"
        assert not report["findings"], f"на чистой колоде выдуманы находки: {report['findings']}"
        print("ok  1. чистая колода проходит, находок нет")

        broken = work / "broken.pptx"
        blank_deck(broken, broken=True)
        code, report = run(broken)
        kinds = {item["kind"] for item in report["findings"]}
        assert code == 1, "испорченная колода прошла как чистая"
        for kind in ("шрифт", "за краем", "заглушка", "числа"):
            assert kind in kinds, f"ворота не заметили «{kind}»; нашли только {kinds}"
        print(f"ok  2. испорченная колода не проходит: {sorted(kinds)}")

        numbers = next(item for item in report["findings"] if item["kind"] == "числа")
        assert "250" in numbers["text"] and "100" in numbers["text"], \
            f"в находке про числа нет самих чисел: {numbers['text']}"
        print("ok  3. арифметика названа числами: заявлено против суммы слагаемых")

        edge = next(item for item in report["findings"] if item["kind"] == "за краем")
        assert "слева" in edge["text"], f"не сказано, куда уехала фигура: {edge['text']}"
        print("ok  4. сказано, куда именно вышла фигура")

        tiny = next(item for item in report["findings"] if item["kind"] == "шрифт")
        assert "6" in tiny["text"], f"в находке про шрифт нет размера: {tiny['text']}"
        print("ok  5. назван размер шрифта")

        print("\nPASS — ворота проверены")
        return 0
    finally:
        shutil.rmtree(work, ignore_errors=True)


if __name__ == "__main__":
    raise SystemExit(main())
