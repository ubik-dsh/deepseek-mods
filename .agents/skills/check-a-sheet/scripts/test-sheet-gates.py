"""Проверка ворот для таблиц: чистая книга проходит, испорченная — ловится.

Книги собираются кодом, потом прогоняются через LibreOffice: он пересчитывает
формулы и кладёт в файл посчитанные значения. Без этого шага openpyxl видит
только текст формул, и проверять «сошлось ли» было бы не на чем.

    python test-sheet-gates.py
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
SCRIPT = HERE / "sheet_gates.py"
SOFFICE = shutil.which("soffice") or r"C:\Program Files\LibreOffice\program\soffice.exe"


def blank_book(path: Path, *, broken: bool) -> None:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Продажи"
    sheet["A1"], sheet["B1"], sheet["C1"] = "товар", "цена", "количество"
    rows = [("ручка", 100, 2), ("тетрадь", 50, 4), ("линейка", 30, 5)]
    for index, (name, price, count) in enumerate(rows, start=2):
        sheet.cell(row=index, column=1, value=name)
        sheet.cell(row=index, column=2, value=price)
        sheet.cell(row=index, column=3, value=count)
        sheet.cell(row=index, column=4, value=f"=B{index}*C{index}")
    sheet["A6"] = "Итого"
    sheet["B6"] = "=SUM(B2:B4)"
    sheet["C6"] = "=SUM(C2:C4)"
    sheet["D6"] = "=SUM(D2:D4)"

    if broken:
        sheet["B7"] = "=1/0"                       # ошибка формулы
        sheet["B8"] = "1 234,5"                     # число, записанное текстом
        sheet["C8"] = 999                           # разрыв столбца: тут формулы, тут константа
        sheet["D9"] = 777                           # константа в столбце формул
        sheet["A3"] = "ручка"                       # дубликат ключа в первом столбце
        sheet["D6"] = 9999                          # «Итого» не сходится
        sheet.merge_cells("A10:B11")                # объединение внутри строк данных
        sheet["A10"] = "сводка"
        sheet["C11"] = "рядом с объединением"        # заполнено ВНЕ диапазона — вот что делает это дефектом

    book.save(str(path))


def recalc(path: Path, workdir: Path) -> Path:
    """Прогнать через LibreOffice: формулы пересчитаются, значения лягут в файл."""
    result = subprocess.run([SOFFICE, "--headless", "--convert-to", "xlsx",
                             "--outdir", str(workdir / "recalc"), str(path)],
                            capture_output=True, text=True, timeout=300)
    target = workdir / "recalc" / path.name
    if result.returncode != 0 or not target.exists():
        raise SystemExit(f"LibreOffice не пересчитал: {result.stderr[:200]}")
    return target


def run(book: Path) -> tuple[int, dict]:
    result = subprocess.run([sys.executable, str(SCRIPT), str(book), "--json"],
                            capture_output=True, text=True, encoding="utf-8")
    if not result.stdout.strip():
        raise SystemExit(f"ворота молчат, код {result.returncode}: {result.stderr[:300]}")
    return result.returncode, json.loads(result.stdout)


def main() -> int:
    work = Path(tempfile.mkdtemp(prefix="sheet-gates-test-"))
    try:
        clean = work / "clean.xlsx"
        blank_book(clean, broken=False)
        clean_recalc = recalc(clean, work)
        code, report = run(clean_recalc)
        assert code == 0, f"чистая книга не прошла: {report['findings']}"
        assert not report["findings"], f"на чистой книге выдуманы находки: {report['findings']}"
        assert report["formula_cells"] >= 6, f"формул найдено {report['formula_cells']}, ждали не меньше 6"
        print(f"ok  1. чистая книга проходит; формул найдено {report['formula_cells']}")

        broken = work / "broken.xlsx"
        blank_book(broken, broken=True)
        broken_recalc = recalc(broken, work)
        code, report = run(broken_recalc)
        kinds = {item["kind"] for item in report["findings"]}
        assert code == 1, "испорченная книга прошла как чистая"
        for kind in ("ошибка формулы", "итоги", "объединённые"):
            assert kind in kinds, f"ворота не заметили «{kind}»; нашли {kinds}"
        print(f"ok  2. испорченная книга не проходит: {sorted(kinds)}")

        totals = next(item for item in report["findings"] if item["kind"] == "итоги")
        assert "9999" in totals["text"], f"в находке про итоги нет заявленного числа: {totals['text']}"
        print("ok  3. итоги названы числами: заявлено против суммы слагаемых")

        warns = {item["kind"] for item in report["warnings"]}
        for kind in ("число как текст", "разрыв столбца", "дубликат ключа"):
            assert kind in warns, f"среди замечаний нет «{kind}»: {warns}"
        print(f"ok  4. замечания на месте: {sorted(warns)}")

        numeric_text = next(item for item in report["warnings"] if item["kind"] == "число как текст")
        assert "1 234,5" in numeric_text["text"], f"не названо само значение: {numeric_text['text']}"
        print("ok  5. число-как-текст показано целиком")

        print("\nPASS — ворота для таблиц проверены")
        return 0
    finally:
        shutil.rmtree(work, ignore_errors=True)


if __name__ == "__main__":
    raise SystemExit(main())
